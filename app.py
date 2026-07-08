import streamlit as st
import pandas as pd
from io import BytesIO
from openpyxl.styles import Font
from calculator import run_calculation, generate_config

st.set_page_config(page_title="门店提成一键计算工具", layout="centered")
st.title("门店提成一键计算系统")
st.caption("自动合并多备用金Sheet LS，完整复刻Excel绩效提成公式")
st.divider()

# 文件上传区域
col1, col2 = st.columns(2)
with col1:
    business_upload = st.file_uploader("① 业务总表xlsx", type=["xlsx"], help="包含完成情况、销售明细、所有备用金Sheet")
with col2:
    staff_upload = st.file_uploader("② 人员配置表xlsx", type=["xlsx"])

# 模板下载
st.subheader("人员配置模板下载")
temp_df = generate_config()
buf_temp = BytesIO()
with pd.ExcelWriter(buf_temp, engine="openpyxl") as writer:
    temp_df.to_excel(writer, index=False, sheet_name="配置模板")
st.download_button(
    label="下载人员配置模板",
    data=buf_temp.getvalue(),
    file_name="人员配置模板.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
st.divider()

# 计算逻辑
can_run = business_upload is not None and staff_upload is not None
if st.button("开始计算", type="primary", disabled=not can_run):
    if business_upload is None or staff_upload is None:
        st.error("请先上传【业务总表】和【人员配置表】两个Excel文件")
    else:
        try:
            with st.spinner("正在计算，请稍候..."):
                # 固定接收3个返回值，和calculator严格对应
                df_res, df_store_sum, df_sales_detail = run_calculation(
                    business_upload,
                    staff_upload,
                    store_sheet="26.05完成情况",
                    sales_sheet="26.05销售明细"
                )
            st.success("计算完成")

            # 统计卡片
            t1, t2, t3 = st.columns(3)
            t1.metric("核算总人数", len(df_res))
            t2.metric("绩效合计", f"¥{df_res['绩效小计'].sum():,.2f}")
            t3.metric("提成合计", f"¥{df_res['提成总计'].sum():,.2f}")

            st.dataframe(df_res, hide_index=True, use_container_width=True)

            # 导出Excel
            output_buf = BytesIO()
            with pd.ExcelWriter(output_buf, engine="openpyxl") as writer:
                df_res.to_excel(writer, index=False, sheet_name="人员提成结果")
                df_store_sum.to_excel(writer, index=False, sheet_name="门店基数汇总")
                df_sales_detail.to_excel(writer, index=False, sheet_name="销售明细(含Z列)")

                # 备注非空的姓名标蓝
                wb = writer.book
                ws = wb["人员提成结果"]
                blue_font = Font(color="FF0000FF", bold=True)
                headers = [cell.value for cell in ws[1]]
                remark_idx = None
                for idx, h in enumerate(headers):
                    if h == "备注":
                        remark_idx = idx + 1
                        break
                if remark_idx:
                    for row_num in range(2, len(df_res) + 2):
                        remark_cell = ws.cell(row=row_num, column=remark_idx)
                        if remark_cell.value and str(remark_cell.value).strip():
                            ws.cell(row=row_num, column=1).font = blue_font

            st.download_button(
                "下载完整结果Excel",
                data=output_buf.getvalue(),
                file_name="门店提成计算结果.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
        except Exception as err:
            st.error(f"计算失败：{str(err)}")
            st.info("排查方向：1、Sheet名称不匹配；2、配置表必填列缺失；3、表格存在非法格式")

# 使用说明
with st.expander("使用说明"):
    st.markdown("""
1. 业务文件自动识别所有名称含「备用金」的工作表，合并统计LS扣款
2. 产品类型自动根据品牌匹配，无需手动填写
3. 销售Y/Z列完整复刻原Excel全部特殊订单规则
4. 人员备注含病假/整月缺勤自动清零绩效提成，导出姓名标蓝
5. 配置表新增「个人提成调整」用于手工补差扣款
6. 输出三张工作表：人员薪酬、门店汇总、原始销售明细核对
""")
